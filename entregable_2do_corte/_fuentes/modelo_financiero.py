"""Modelo financiero de Hand Sing Kids (fuente única para documento, póster y .xlsx).

Formulación (GUIA_CALCULO_Y_JUSTIFICACION_ROI.md):
    Beneficio neto anual = Beneficios brutos - OPEX
    ROI (%)              = Beneficio neto anual / CAPEX x 100
    Payback (meses)      = CAPEX / (Beneficio neto anual / 12)

Unidad de análisis: UN aula de música de una institución educativa de Montería o
Cereté, porque los datos de campo disponibles se reportan por aula. Procedencia de
los insumos (columna «Tipo» de la hoja):
  Guía UPB   tarifa de ingeniería dentro del rango de la guía.
  PINT2      valor reportado por el estudio de campo del Equipo 13 (Proyecto Integrador II,
             2026). Es una fuente secundaria interna: este equipo no lo verificó.
  SUPUESTO   estimación propia, sin respaldo de campo.
  Verificable comprobado en el código (p. ej. no hay nube ni API de IA).
`python modelo_financiero.py` regenera modelo_financiero.json, el .xlsx y la figura.
"""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parents[1]

P = dict(
    tarifa_ing=30_000,            # Guía UPB: rango 25.000-40.000 COP/h
    horas_dev_pint2=120,          # PINT2: horas de programación y optimización reportadas
    horas_dev_corte2=80,          # SUPUESTO: gobernanza, pruebas, verificación y documentación de este corte
    hardware_kit=300_000,         # PINT2: cámara web 720p y periféricos de prueba
    prototipado=150_000,          # PINT2: soportes y materiales
    configuracion=200_000,        # PINT2: entorno, herramientas y puesta a punto (el software es libre)
    licencias=0,                  # Verificable: dependencias de código abierto
    horas_capac=4, tarifa_docente=30_000,   # PINT2: 4 h de entrenamiento a docentes a 30.000 COP/h
    consulta_juridica=600_000,    # SUPUESTO: Ley 1581 de 2012 (datos de menores)
    mantenimiento=300_000,        # PINT2: soporte preventivo, corrección de errores y librerías
    pct_reposicion=0.10,          # SUPUESTO: reposición anual de hardware
    energia=60_000,               # SUPUESTO
    nube=0, api_ia=0, conectividad=0,   # Verificable: el código no usa nube, API de IA ni red
    horas_nivelacion_sem=1.0, semanas=40,   # PINT2: 1,0 a 1,5 h/sem de nivelación individual (se usa el mínimo)
    papeleria_anual=600_000,      # PINT2: fotocopias y cuadernos por aula
    instrumental_evitado=1_000_000,   # PINT2: instrumentos físicos no adquiridos (solo escenario optimista)
    aulas_institucion=4,          # SUPUESTO: aulas de música que adoptan el sistema en una institución
)
FACTOR_CONSERVADOR = 0.75
N_INSTITUCIONES_ESCALA = 3
ESC_CONS, ESC_BASE = "Conservador (beneficios -25 %)", "Base (1 aula)"
ESC_OPT, ESC_ESCALA = "Optimista (con instrumental evitado)", "Escala (3 instituciones comparten el desarrollo)"
ESC_MAX = "Escala y optimista"
ESC_INST = "Institución con 4 aulas"


def calcular(p: dict = P) -> dict:
    horas = p["horas_dev_pint2"] + p["horas_dev_corte2"]
    hw = p["hardware_kit"] + p["prototipado"]
    capex = {
        "Horas de ingeniería y desarrollo": horas * p["tarifa_ing"],
        "Cámara web, periféricos y soportes": hw,
        "Licenciamiento y configuración inicial": p["licencias"] + p["configuracion"],
        "Capacitación a docentes": p["horas_capac"] * p["tarifa_docente"],
        "Consulta jurídica de protección de datos": p["consulta_juridica"],
    }
    opex = {
        "Mantenimiento y soporte técnico": p["mantenimiento"],
        "Reposición de hardware": p["pct_reposicion"] * hw,
        "Energía eléctrica": p["energia"],
        "Nube, API de IA y conectividad": p["nube"] + p["api_ia"] + p["conectividad"],
    }
    benef = {
        "Ahorro de tiempo docente en nivelación individual": p["horas_nivelacion_sem"] * p["semanas"] * p["tarifa_docente"],
        "Reducción del gasto en papelería didáctica": p["papeleria_anual"],
    }
    C, O, B = sum(capex.values()), sum(opex.values()), sum(benef.values())
    inst = p["instrumental_evitado"]

    def kpi(c, o, b):
        neto = b - o
        return {"capex": c, "opex": o, "beneficio_bruto": b, "beneficio_neto": neto, "roi_pct": neto / c * 100,
                "payback_meses": c / (neto / 12) if neto > 0 else None, "beneficio_neto_mensual": neto / 12}

    dev = capex["Horas de ingeniería y desarrollo"]
    c_esc = C - dev + dev / N_INSTITUCIONES_ESCALA
    na = p["aulas_institucion"]
    c_inst = (dev + na * hw + p["licencias"] + p["configuracion"] + na * p["horas_capac"] * p["tarifa_docente"]
              + p["consulta_juridica"])
    o_inst = p["mantenimiento"] + na * (p["pct_reposicion"] * hw + p["energia"])
    escenarios = {
        ESC_CONS: kpi(C, O, B * FACTOR_CONSERVADOR),
        ESC_BASE: kpi(C, O, B),
        ESC_OPT: kpi(C, O, B + inst),
        ESC_INST: kpi(c_inst, o_inst, na * B),
        ESC_ESCALA: kpi(c_esc, O, B),
        ESC_MAX: kpi(c_esc, O, B + inst),
    }
    return {"parametros": p, "capex": capex, "opex": opex, "beneficios": benef, "instrumental": inst,
            "horas_dev": horas, "escenarios": escenarios, "totales": {"capex": C, "opex": O, "beneficios": B}}


def cop(v: float) -> str:
    return "$ " + f"{v:,.0f}".replace(",", ".")


def escribir_xlsx(res: dict, path: Path) -> None:
    wb = Workbook()
    hfill, yfill, gfill = (PatternFill("solid", fgColor=c) for c in ("141414", "FFCC00", "F5F5F5"))
    thin = Side(style="thin", color="BBBBBB")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)

    def head(ws, row, labels):
        for i, t in enumerate(labels, 1):
            c = ws.cell(row=row, column=i, value=t)
            c.font, c.fill, c.border = Font(bold=True, color="FFFFFF"), hfill, box
            c.alignment = Alignment(wrap_text=True, vertical="center")

    ws = wb.active
    ws.title = "Supuestos"
    ws["A1"] = "Hand Sing Kids — supuestos del modelo financiero (COP, año 1, una aula de música)"
    ws["A1"].font = Font(bold=True, size=13)
    head(ws, 3, ["Parámetro", "Valor", "Unidad", "Tipo", "Nota"])
    filas = [
        ("tarifa_ing", "Tarifa de ingeniería junior", "COP/h", "Guía UPB", "Rango de la guía: 25.000 a 40.000"),
        ("horas_dev_pint2", "Horas de desarrollo reportadas por PINT2", "h", "PINT2", "Fuente secundaria interna, no verificada"),
        ("horas_dev_corte2", "Horas de este corte (gobernanza, pruebas, documentación)", "h", "SUPUESTO", "Validar con la bitácora del equipo"),
        ("hardware_kit", "Cámara web 720p y periféricos", "COP", "PINT2", "Un kit por aula"),
        ("prototipado", "Soportes y materiales", "COP", "PINT2", ""),
        ("configuracion", "Configuración inicial y entorno", "COP", "PINT2", "El software es de código abierto"),
        ("licencias", "Licenciamiento de software", "COP", "Verificable", "Dependencias de código abierto"),
        ("horas_capac", "Horas de capacitación a docentes", "h", "PINT2", ""),
        ("tarifa_docente", "Costo hora docente", "COP/h", "PINT2", "Validar con la nómina de la institución"),
        ("consulta_juridica", "Consulta jurídica (Ley 1581 de 2012)", "COP", "SUPUESTO", "Datos de menores"),
        ("mantenimiento", "Mantenimiento y soporte anual", "COP", "PINT2", ""),
        ("pct_reposicion", "Reposición anual de hardware", "fracción", "SUPUESTO", "Sobre el costo del hardware"),
        ("energia", "Energía eléctrica anual", "COP", "SUPUESTO", ""),
        ("nube", "Infraestructura en la nube", "COP", "Verificable", "El código no usa nube"),
        ("api_ia", "API de inteligencia artificial", "COP", "Verificable", "El código no llama a ninguna API de IA"),
        ("conectividad", "Conectividad", "COP", "Verificable", "La aplicación funciona sin internet"),
        ("horas_nivelacion_sem", "Horas docentes de nivelación individual por semana", "h", "PINT2", "Reportado: 1,0 a 1,5 h; se usa el mínimo; medir en piloto"),
        ("semanas", "Semanas lectivas", "sem", "PINT2", ""),
        ("papeleria_anual", "Gasto anual en papelería por aula", "COP", "PINT2", "Fotocopias y cuadernos de caligrafía musical"),
        ("instrumental_evitado", "Instrumental no adquirido (solo escenario optimista)", "COP", "PINT2", "Beneficio discutible: no se incluye en el escenario base"),
        ("aulas_institucion", "Aulas de música que adoptan el sistema en una institución", "u", "SUPUESTO", "Solo escenario «Institución con 4 aulas»"),
    ]
    ref = {}
    for i, (k, nombre, unidad, tipo, nota) in enumerate(filas, 4):
        ws.cell(row=i, column=1, value=nombre)
        c = ws.cell(row=i, column=2, value=res["parametros"][k])
        c.fill = yfill if tipo in ("SUPUESTO", "PINT2") else gfill
        ws.cell(row=i, column=3, value=unidad); ws.cell(row=i, column=4, value=tipo); ws.cell(row=i, column=5, value=nota)
        ref[k] = f"Supuestos!$B${i}"
        for col in range(1, 6):
            ws.cell(row=i, column=col).border = box
    n = len(filas) + 4
    ws.cell(row=n + 1, column=1, value="Factor de beneficios, escenario conservador"); ws.cell(row=n + 1, column=2, value=FACTOR_CONSERVADOR).fill = yfill
    ws.cell(row=n + 2, column=1, value="Instituciones que comparten el desarrollo (escala)"); ws.cell(row=n + 2, column=2, value=N_INSTITUCIONES_ESCALA).fill = yfill
    ref["fc"], ref["nesc"] = f"Supuestos!$B${n + 1}", f"Supuestos!$B${n + 2}"
    for col, w in zip("ABCDE", (58, 14, 10, 12, 70)):
        ws.column_dimensions[col].width = w

    r = ref
    m = wb.create_sheet("Modelo")
    m["A1"] = "Modelo financiero — escenario Base (un aula de música)"
    m["A1"].font = Font(bold=True, size=13)
    head(m, 3, ["Categoría", "Concepto", "Fórmula (detalle)", "Valor anual (COP)"])
    hw = f"({r['hardware_kit']}+{r['prototipado']})"
    bloques = [
        ("CAPEX", [("Horas de ingeniería y desarrollo", "(horas PINT2 + horas del corte) × tarifa", f"=({r['horas_dev_pint2']}+{r['horas_dev_corte2']})*{r['tarifa_ing']}"),
                   ("Cámara web, periféricos y soportes", "kit + prototipado", f"={hw}"),
                   ("Licenciamiento y configuración inicial", "licencias + configuración", f"={r['licencias']}+{r['configuracion']}"),
                   ("Capacitación a docentes", "horas × tarifa docente", f"={r['horas_capac']}*{r['tarifa_docente']}"),
                   ("Consulta jurídica de protección de datos", "", f"={r['consulta_juridica']}")], "TOTAL CAPEX (I0)"),
        ("OPEX", [("Mantenimiento y soporte técnico", "anual", f"={r['mantenimiento']}"),
                  ("Reposición de hardware", "% × hardware", f"={r['pct_reposicion']}*{hw}"),
                  ("Energía eléctrica", "", f"={r['energia']}"),
                  ("Nube, API de IA y conectividad", "el código no los usa", f"={r['nube']}+{r['api_ia']}+{r['conectividad']}")], "TOTAL OPEX anual"),
        ("BENEFICIO", [("Ahorro de tiempo docente en nivelación individual", "h/sem × semanas × costo hora", f"={r['horas_nivelacion_sem']}*{r['semanas']}*{r['tarifa_docente']}"),
                       ("Reducción del gasto en papelería didáctica", "por aula y año", f"={r['papeleria_anual']}")], "TOTAL BENEFICIOS BRUTOS anuales"),
    ]
    row = 4
    tot = {}
    for cat, lineas, etiqueta in bloques:
        r0 = row
        for con, det, f in lineas:
            for col, v in enumerate((cat, con, det, f), 1):
                m.cell(row=row, column=col, value=v).border = box
            row += 1
        m.cell(row=row, column=2, value=etiqueta).font = Font(bold=True)
        m.cell(row=row, column=4, value=f"=SUM(D{r0}:D{row - 1})").font = Font(bold=True)
        tot[cat] = row
        row += 2
    cx, ox, bx = tot["CAPEX"], tot["OPEX"], tot["BENEFICIO"]
    neto = row
    for nombre, f, fmt in (("Beneficio neto anual", f"=D{bx}-D{ox}", '"$ "#,##0'),
                           ("ROI calculado (año 1), %", f"=D{neto}/D{cx}*100", "0.0"),
                           ("Beneficio neto mensual", f"=D{neto}/12", '"$ "#,##0'),
                           ("Periodo de retorno (payback), meses", f"=D{cx}/D{neto + 2}", "0.0")):
        m.cell(row=row, column=2, value=nombre).font = Font(bold=True)
        c = m.cell(row=row, column=4, value=f)
        c.font, c.fill, c.number_format = Font(bold=True), yfill, fmt
        row += 1
    for rr in range(4, neto):
        m.cell(row=rr, column=4).number_format = '"$ "#,##0'
    for col, w in zip("ABCD", (12, 52, 52, 22)):
        m.column_dimensions[col].width = w

    e = wb.create_sheet("Escenarios")
    e["A1"] = "Escenarios frente a los criterios de la guía"
    e["A1"].font = Font(bold=True, size=13)
    head(e, 3, ["Escenario", "CAPEX", "OPEX anual", "Beneficios brutos", "Beneficio neto", "ROI año 1 (%)", "Payback (meses)", "¿ROI ≥ 30 %?", "¿Payback ≤ 12 m?"])
    dev = f"({r['horas_dev_pint2']}+{r['horas_dev_corte2']})*{r['tarifa_ing']}"
    capex_esc = f"=Modelo!D{cx}-{dev}+{dev}/{r['nesc']}"
    filas_e = [(ESC_CONS, f"=Modelo!D{cx}", f"=Modelo!D{bx}*{r['fc']}"),
               (ESC_BASE, f"=Modelo!D{cx}", f"=Modelo!D{bx}"),
               (ESC_OPT, f"=Modelo!D{cx}", f"=Modelo!D{bx}+{r['instrumental_evitado']}"),
               (ESC_INST, f"={dev}+{r['aulas_institucion']}*{hw}+{r['licencias']}+{r['configuracion']}+{r['aulas_institucion']}*{r['horas_capac']}*{r['tarifa_docente']}+{r['consulta_juridica']}",
                f"={r['aulas_institucion']}*Modelo!D{bx}"),
               (ESC_ESCALA, capex_esc, f"=Modelo!D{bx}"),
               (ESC_MAX, capex_esc, f"=Modelo!D{bx}+{r['instrumental_evitado']}")]
    for i, (nm, cxf, bxf) in enumerate(filas_e, 4):
        e.cell(row=i, column=1, value=nm)
        ox_f = (f"={r['mantenimiento']}+{r['aulas_institucion']}*({r['pct_reposicion']}*{hw}+{r['energia']})" if nm == ESC_INST
                else f"=Modelo!D{ox}")
        e.cell(row=i, column=2, value=cxf); e.cell(row=i, column=3, value=ox_f); e.cell(row=i, column=4, value=bxf)
        e.cell(row=i, column=5, value=f"=D{i}-C{i}")
        e.cell(row=i, column=6, value=f"=E{i}/B{i}*100").number_format = "0.0"
        e.cell(row=i, column=7, value=f"=B{i}/(E{i}/12)").number_format = "0.0"
        e.cell(row=i, column=8, value=f'=IF(F{i}>=30,"Sí","No")')
        e.cell(row=i, column=9, value=f'=IF(G{i}<=12,"Sí","No")')
        for col in range(1, 10):
            e.cell(row=i, column=col).border = box
        for col in range(2, 6):
            e.cell(row=i, column=col).number_format = '"$ "#,##0'
    for col, w in zip("ABCDEFGHI", (50, 16, 16, 18, 16, 14, 16, 14, 16)):
        e.column_dimensions[col].width = w

    f = wb.create_sheet("Flujo acumulado")
    f["A1"] = "Posición acumulada = −CAPEX + mes × beneficio neto mensual"
    f["A1"].font = Font(bold=True, size=13)
    head(f, 3, ["Mes", "Conservador", "Base", "Optimista", "Institución 4 aulas", "Escala", "Escala y optimista"])
    for mes in range(0, 61):
        rr = 4 + mes
        f.cell(row=rr, column=1, value=mes)
        for j, er in enumerate((4, 5, 6, 7, 8, 9), 2):
            f.cell(row=rr, column=j, value=f"=-Escenarios!$B${er}+A{rr}*Escenarios!$E${er}/12").number_format = '"$ "#,##0'
    for col in "ABCDEFG":
        f.column_dimensions[col].width = 18
    wb.save(path)


def figura_flujo(res: dict) -> str:
    from svglib import SVG, INK, GRAY_D, TEAL, RED
    # lienzo de 720 para página vertical (6,5 in): 13,5 unidades ≈ 8,8 pt
    W, H = 720, 520
    FS = 13.5
    s = SVG(W, H, "Flujo de caja acumulado por escenario")
    L, Rr, T_, B_ = 70, 705, 14, 330
    esc = res["escenarios"]
    meses = 60
    pos = lambda e, m: -e["capex"] + m * e["beneficio_neto_mensual"]
    allv = [pos(e, m) for e in esc.values() for m in (0, meses)]
    lo, hi = min(allv) * 1.05, max(allv) * 1.05
    X = lambda m: L + (Rr - L) * m / meses
    Y = lambda v: B_ - (B_ - T_) * (v - lo) / (hi - lo)
    step = 5_000_000
    v = (int(lo // step)) * step
    while v <= hi:
        y = Y(v)
        s.line(L, y, Rr, y, stroke="#E3E3E3" if v else "#888", sw=1 if v else 1.6)
        s.text(L - 8, y, f"{v / 1e6:.0f}", size=FS, anchor="end", weight=500, fill=GRAY_D)
        v += step
    for m in range(0, meses + 1, 12):
        s.line(X(m), B_, X(m), B_ + 6, sw=1.2)
        s.text(X(m), B_ + 20, str(m), size=FS, weight=500, fill=GRAY_D)
    s.text((L + Rr) / 2, B_ + 44, "Meses desde la puesta en marcha", size=FS, weight=600)
    s.text(16, (T_ + B_) / 2, "Posición acumulada (millones de COP)", size=FS, weight=600, rotate=-90)
    s.line(X(12), T_, X(12), B_, stroke=RED, sw=1.4, dash="6 4")
    s.text(X(12) + 6, T_ + 10, "Límite de la guía: 12 meses", size=FS, anchor="start", weight=600, fill=RED)
    estilos = {ESC_CONS: (GRAY_D, "7 4"), ESC_BASE: (INK, None), ESC_OPT: ("#B58900", "3 3"), ESC_ESCALA: (TEAL, None), ESC_MAX: ("#0B6B61", "10 4"), ESC_INST: ("#7A3E9D", None)}
    for nm, e in esc.items():
        col, da = estilos[nm]
        s.path([(X(0), Y(pos(e, 0))), (X(meses), Y(pos(e, meses)))], stroke=col, sw=3.4 if nm == ESC_BASE else 2.6, dash=da, end=None)
        pb = e["payback_meses"]
        if pb and pb <= meses:
            s.add(f'<circle cx="{X(pb):.1f}" cy="{Y(0):.1f}" r="6" fill="{col}" stroke="#fff" stroke-width="2"/>')
    ly = B_ + 78
    for i, (nm, e) in enumerate(esc.items()):
        col, da = estilos[nm]
        x = 10 + (i % 2) * 360
        y = ly + (i // 2) * 40
        d = f' stroke-dasharray="{da}"' if da else ""
        s.add(f'<line x1="{x}" y1="{y}" x2="{x + 34}" y2="{y}" stroke="{col}" stroke-width="3"{d}/>')
        s.text(x + 42, y - 8, nm.split(" (")[0], size=FS, anchor="start", weight=700)
        pb = f'{e["payback_meses"]:.1f}'.replace(".", ",")
        s.text(x + 42, y + 9, f'ROI {e["roi_pct"]:.1f} %'.replace(".", ",") + f"; payback {pb} meses", size=FS, anchor="start", weight=500, fill=GRAY_D)
    return s.svg()


if __name__ == "__main__":
    import sys
    sys.path.insert(0, str(Path(__file__).parent))
    res = calcular()
    (Path(__file__).parent / "modelo_financiero.json").write_text(json.dumps(res, ensure_ascii=False, indent=1), encoding="utf-8")
    escribir_xlsx(res, ROOT / "entrega_teams" / "04_Modelo_Financiero_ROI_Equipo_XX.xlsx")
    (ROOT / "recursos" / "diagramas" / "flujo_caja_acumulado.svg").write_text(figura_flujo(res), encoding="utf-8")
    t = res["totales"]
    print("CAPEX", cop(t["capex"]), "| OPEX", cop(t["opex"]), "| Beneficios", cop(t["beneficios"]))
    for k, e in res["escenarios"].items():
        print(f"{k:55s} neto {cop(e['beneficio_neto'])}  ROI {e['roi_pct']:.1f}%  payback {e['payback_meses']:.1f} m")
