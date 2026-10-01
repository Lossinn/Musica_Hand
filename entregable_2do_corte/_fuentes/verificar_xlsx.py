"""Evalúa de forma independiente las fórmulas del .xlsx y las compara con modelo_financiero.json."""
import json
import re
from pathlib import Path

from openpyxl import load_workbook

F = Path(__file__).parent
wb = load_workbook(F.parent / "entrega_teams" / "04_Modelo_Financiero_ROI_Equipo_XX.xlsx")
vals = {}


def val(sh, ref):
    k = (sh, ref.replace("$", ""))
    if k in vals:
        return vals[k]
    c = wb[sh][k[1]].value
    v = ev(sh, c[1:]) if isinstance(c, str) and c.startswith("=") else c
    vals[k] = v
    return v


def ev(sh, expr):
    expr = re.sub(r"SUM\(([A-Z]+)(\d+):([A-Z]+)(\d+)\)",
                  lambda m: "(" + "+".join(f"{m.group(1)}{r}" for r in range(int(m.group(2)), int(m.group(4)) + 1)) + ")", expr)

    def rep(m):
        s = (m.group(1) or sh + "!")[:-1]
        return repr(val(s, m.group(2)))
    return eval(re.sub(r"((?:[A-Za-zÁÉÍÓÚáéíóú ]+)!)?(\$?[A-Z]{1,2}\$?\d+)", rep, expr))


res = json.loads((F / "modelo_financiero.json").read_text(encoding="utf-8"))
e = wb["Escenarios"]
ok = True
for r in range(4, 10):
    nombre = e.cell(row=r, column=1).value
    x = [val("Escenarios", f"{c}{r}") for c in "BCDEFG"]
    j = res["escenarios"][nombre]
    esperado = [j["capex"], j["opex"], j["beneficio_bruto"], j["beneficio_neto"], j["roi_pct"], j["payback_meses"]]
    coincide = all(abs(a - b) < 1e-6 for a, b in zip(x, esperado))
    ok &= coincide
    print(("OK   " if coincide else "DIFIERE"), nombre, [round(v, 1) for v in x])
print("TODAS COINCIDEN" if ok else "HAY DIFERENCIAS")
